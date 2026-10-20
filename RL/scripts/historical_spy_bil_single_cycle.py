#!/usr/bin/env python3
"""Contrôle un unique cycle historique SPY/BIL préspécifié.

La frontière ne choisit aucun poids, ne parcourt aucun historique et ne calcule
aucune performance. Elle vérifie les sources, prépare l'intention aux clôtures,
puis seulement lit les ouvertures et transmet le lot au registre strict.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from enum import Enum
from pathlib import Path
from typing import Any, Mapping
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from synthetic_decision_adapter import (
    AdapterError,
    DatedPrice,
    DecisionIntent,
    ExecutionOutcome,
    ExecutionStatus,
    PriceStatus,
    SyntheticDecisionAdapter,
)
from synthetic_portfolio_ledger import (
    AccountingSnapshot,
    LedgerError,
    SyntheticPortfolioLedger,
    SyntheticSettlementCalendar,
    SyntheticTradingCalendar,
)


SCHEMA_VERSION = "historical_spy_bil_single_cycle/v1"
PROVENANCE = "HISTORICAL_AUDITED_PROXY"
SYMBOLS = ("SPY", "BIL")
SHA256_PATTERN = re.compile(r"^[0-9a-f]{64}$")
DECIMAL_PATTERN = re.compile(r"^(?:0|[1-9]\d*)(?:\.\d+)?$")


class HistoricalBoundaryError(ValueError):
    """Contrat ou intégrité incompatible avec la frontière mono-cycle."""


class DocumentaryBlock(HistoricalBoundaryError):
    """Preuve documentaire insuffisante avant toute mutation financière."""


class HistoricalCycleVerdict(str, Enum):
    VALIDATED = "HISTORICAL_CYCLE_VALIDATED"
    NOT_EVALUABLE = "HISTORICAL_CYCLE_NOT_EVALUABLE"
    FAILED = "HISTORICAL_CYCLE_FAILED"


@dataclass(frozen=True)
class VerifiedSource:
    identifier: str
    path: Path
    sha256: str
    record: Mapping[str, Any]


@dataclass(frozen=True)
class CycleContract:
    config_sha256: str
    cycle_id: str
    initial_cash: Decimal
    target_weight_spy: Decimal
    decision_date: date
    decision_at: datetime
    trade_date: date
    settlement_date: date
    trading_dates: tuple[date, ...]
    settlement_dates: tuple[date, ...]
    fees: Mapping[str, Decimal]
    source_ids: Mapping[str, str]
    event_start: date
    event_end: date


@dataclass
class PreparedHistoricalCycle:
    contract: CycleContract
    sources: Mapping[str, VerifiedSource]
    source_documents: Mapping[str, Mapping[str, Any]]
    ledger: SyntheticPortfolioLedger
    adapter: SyntheticDecisionAdapter
    intent: DecisionIntent
    closes: Mapping[str, Decimal]
    trace: list[str]


@dataclass(frozen=True)
class HistoricalCycleResult:
    verdict: HistoricalCycleVerdict
    reason: str
    cycle_id: str | None = None
    config_sha256: str | None = None
    trace: tuple[str, ...] = ()
    source_hashes: Mapping[str, str] | None = None
    closes: Mapping[str, Decimal] | None = None
    opens: Mapping[str, Decimal] | None = None
    intent: DecisionIntent | None = None
    outcome: ExecutionOutcome | None = None
    before_settlement: AccountingSnapshot | None = None
    after_settlement: AccountingSnapshot | None = None
    settled_ids: tuple[int, ...] = ()
    wealth_before_settlement: Decimal | None = None
    wealth_after_settlement: Decimal | None = None

    def to_record(self) -> dict[str, Any]:
        record: dict[str, Any] = {
            "verdict": self.verdict.value,
            "reason": self.reason,
            "cycle_id": self.cycle_id,
            "config_sha256": self.config_sha256,
            "trace": list(self.trace),
            "source_hashes": dict(self.source_hashes or {}),
            "closes": _decimal_mapping(self.closes),
            "opens": _decimal_mapping(self.opens),
            "settled_ids": list(self.settled_ids),
            "wealth_before_settlement": _decimal_or_none(
                self.wealth_before_settlement
            ),
            "wealth_after_settlement": _decimal_or_none(
                self.wealth_after_settlement
            ),
        }
        if self.intent is not None:
            record["intent"] = {
                "intent_id": self.intent.intent_id,
                "content_sha256": self.intent.content_sha256,
                "target_weight_spy": _decimal_text(
                    self.intent.target_weight_spy
                ),
                "target_quantities": dict(self.intent.target_quantities),
                "orders": [
                    {"symbol": item.symbol, "quantity": item.quantity}
                    for item in self.intent.orders
                ],
            }
        if self.outcome is not None:
            record["outcome"] = {
                "status": self.outcome.status.value,
                "reason": self.outcome.reason,
                "settlement_date": (
                    None
                    if self.outcome.settlement_date is None
                    else self.outcome.settlement_date.isoformat()
                ),
                "funding_source": self.outcome.funding_source.value,
            }
        record["before_settlement"] = _snapshot_record(self.before_settlement)
        record["after_settlement"] = _snapshot_record(self.after_settlement)
        return record


def _decimal_text(value: Decimal) -> str:
    if value == 0:
        return "0"
    return format(value.normalize(), "f")


def _decimal_or_none(value: Decimal | None) -> str | None:
    return None if value is None else _decimal_text(value)


def _decimal_mapping(values: Mapping[str, Decimal] | None) -> dict[str, str] | None:
    if values is None:
        return None
    return {key: _decimal_text(value) for key, value in sorted(values.items())}


def _snapshot_record(snapshot: AccountingSnapshot | None) -> dict[str, Any] | None:
    if snapshot is None:
        return None
    return {
        "settled_cash": _decimal_text(snapshot.settled_cash),
        "economic_positions": dict(sorted(snapshot.economic_positions.items())),
        "settled_positions": dict(sorted(snapshot.settled_positions.items())),
        "sale_receivables": _decimal_text(snapshot.sale_receivables),
        "purchase_payables": _decimal_text(snapshot.purchase_payables),
        "distribution_receivables": _decimal_text(
            snapshot.distribution_receivables
        ),
        "fees_recognized": _decimal_text(snapshot.fees_recognized),
        "pending_settlement_count": snapshot.pending_settlement_count,
        "unpaid_distribution_count": snapshot.unpaid_distribution_count,
    }


def _mapping(value: Any, label: str) -> Mapping[str, Any]:
    if not isinstance(value, dict):
        raise HistoricalBoundaryError(f"{label} doit être un objet")
    return value


def _exact_keys(value: Mapping[str, Any], expected: set[str], label: str) -> None:
    actual = set(value)
    if actual != expected:
        raise HistoricalBoundaryError(
            f"{label}: absents={sorted(expected - actual)}, "
            f"inattendus={sorted(actual - expected)}"
        )


def _text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value:
        raise HistoricalBoundaryError(f"{label} doit être une chaîne non vide")
    return value


def _date(value: Any, label: str) -> date:
    text = _text(value, label)
    try:
        parsed = date.fromisoformat(text)
    except ValueError as exc:
        raise HistoricalBoundaryError(f"{label} doit être une date ISO") from exc
    if parsed.isoformat() != text:
        raise HistoricalBoundaryError(f"{label} doit être une date ISO canonique")
    return parsed


def _decimal(value: Any, label: str) -> Decimal:
    text = _text(value, label)
    if not DECIMAL_PATTERN.fullmatch(text):
        raise HistoricalBoundaryError(f"{label} doit être un décimal textuel")
    try:
        parsed = Decimal(text)
    except InvalidOperation as exc:
        raise HistoricalBoundaryError(f"{label} est invalide") from exc
    if not parsed.is_finite() or parsed < 0:
        raise HistoricalBoundaryError(f"{label} doit être fini et non négatif")
    return parsed


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _verified_json(path: Path) -> tuple[Mapping[str, Any], str]:
    try:
        raw = path.read_bytes()
        sidecar = path.with_suffix(path.suffix + ".sha256").read_text(
            encoding="ascii"
        )
    except OSError as exc:
        raise HistoricalBoundaryError("configuration ou SHA-256 inaccessible") from exc
    parts = sidecar.strip().split()
    digest = hashlib.sha256(raw).hexdigest()
    if (
        len(parts) != 2
        or parts[1] != path.name
        or not SHA256_PATTERN.fullmatch(parts[0])
        or digest != parts[0]
    ):
        raise HistoricalBoundaryError("empreinte de configuration incorrecte")
    try:
        document = json.loads(raw.decode("utf-8"), parse_float=Decimal)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HistoricalBoundaryError("configuration JSON invalide") from exc
    return _mapping(document, "configuration"), digest


def load_cycle_contract(path: str | Path) -> CycleContract:
    root, digest = _verified_json(Path(path))
    _exact_keys(
        root,
        {
            "schema_version",
            "cycle_id",
            "provenance",
            "scope",
            "account",
            "decision",
            "execution",
            "sources",
            "calendars",
            "event_window",
        },
        "configuration",
    )
    if root["schema_version"] != SCHEMA_VERSION:
        raise HistoricalBoundaryError("version de schéma inconnue")
    if root["provenance"] != PROVENANCE:
        raise HistoricalBoundaryError("provenance historique invalide")
    if root["scope"] != "SINGLE_PRESPECIFIED_CYCLE":
        raise HistoricalBoundaryError("seul un cycle préspécifié est admis")

    account = _mapping(root["account"], "account")
    _exact_keys(
        account,
        {
            "currency",
            "initial_cash",
            "initial_positions",
            "strict_temporal",
            "funding",
            "whole_shares",
            "external_fees",
            "residual_cash_interest",
        },
        "account",
    )
    positions = _mapping(account["initial_positions"], "initial_positions")
    fees_raw = _mapping(account["external_fees"], "external_fees")
    if set(positions) != set(SYMBOLS) or any(positions[item] != 0 for item in SYMBOLS):
        raise HistoricalBoundaryError("les positions initiales doivent être SPY=BIL=0")
    if set(fees_raw) != set(SYMBOLS):
        raise HistoricalBoundaryError("les frais doivent couvrir SPY et BIL")
    fees = {item: _decimal(fees_raw[item], f"frais {item}") for item in SYMBOLS}
    if (
        account["currency"] != "USD"
        or account["strict_temporal"] is not True
        or account["funding"] != "H_FIN_NET"
        or account["whole_shares"] is not True
        or _decimal(account["residual_cash_interest"], "intérêt du cash") != 0
        or any(fees.values())
    ):
        raise HistoricalBoundaryError("conventions comptables du cycle modifiées")

    decision = _mapping(root["decision"], "decision")
    _exact_keys(
        decision,
        {
            "decision_price_date",
            "decision_at",
            "timezone",
            "decision_price_field",
            "decision_price_status",
            "target_weight_spy",
            "target_weight_bil",
        },
        "decision",
    )
    if decision["decision_price_field"] != "close":
        raise HistoricalBoundaryError("la décision doit utiliser uniquement le close")
    if decision["decision_price_status"] != "MARKET_OBSERVABLE_PROXY":
        raise DocumentaryBlock("statut point-in-time du close non admissible")
    timezone_name = _text(decision["timezone"], "timezone")
    if timezone_name != "America/New_York":
        raise HistoricalBoundaryError("le fuseau doit être America/New_York")
    try:
        decision_at = datetime.fromisoformat(decision["decision_at"])
        if decision_at.tzinfo is None or decision_at.utcoffset() is None:
            raise ValueError("horodatage sans fuseau")
        local_decision = decision_at.astimezone(ZoneInfo(timezone_name))
    except (TypeError, ValueError) as exc:
        raise HistoricalBoundaryError("instant de décision invalide") from exc
    if local_decision.strftime("%H:%M:%S") != "08:00:00":
        raise HistoricalBoundaryError("la décision doit être fixée à 08:00 ET")
    spy_weight = _decimal(decision["target_weight_spy"], "poids SPY")
    bil_weight = _decimal(decision["target_weight_bil"], "poids BIL")
    if spy_weight != Decimal("0.5") or bil_weight != Decimal("0.5"):
        raise HistoricalBoundaryError("les poids imposés doivent rester 50/50")

    execution = _mapping(root["execution"], "execution")
    _exact_keys(
        execution,
        {
            "trade_date",
            "execution_price_field",
            "execution_price_status",
            "proxy_use_authorized_for_cycle",
            "anomaly_status",
            "settlement_date",
            "settlement_cycle",
        },
        "execution",
    )
    if execution["execution_price_field"] != "open":
        raise HistoricalBoundaryError("le prix d'exécution doit être l'open")
    if (
        execution["execution_price_status"] != "MARKET_OBSERVABLE_PROXY"
        or execution["proxy_use_authorized_for_cycle"] is not True
    ):
        raise DocumentaryBlock("proxy d'ouverture non autorisé")
    if execution["anomaly_status"] != "CLEAR":
        raise DocumentaryBlock("prix contesté pour le cycle")
    if execution["settlement_cycle"] != "T+1":
        raise HistoricalBoundaryError("le cycle de règlement doit être T+1")

    decision_date = _date(decision["decision_price_date"], "decision_price_date")
    trade_date = _date(execution["trade_date"], "trade_date")
    settlement_date = _date(execution["settlement_date"], "settlement_date")
    if local_decision.date() != trade_date or not decision_date < trade_date < settlement_date:
        raise HistoricalBoundaryError("chronologie du cycle incohérente")

    calendars = _mapping(root["calendars"], "calendars")
    _exact_keys(
        calendars,
        {"trading", "settlement", "trading_status", "settlement_status"},
        "calendars",
    )
    if (
        calendars["trading_status"] != "VERIFIED_POINT_IN_TIME"
        or calendars["settlement_status"] != "VERIFIED_POINT_IN_TIME"
    ):
        raise DocumentaryBlock("calendrier non vérifié")
    trading_dates = tuple(_date(item, "trading date") for item in calendars["trading"])
    settlement_dates = tuple(
        _date(item, "settlement date") for item in calendars["settlement"]
    )
    if trading_dates != (decision_date, trade_date):
        raise HistoricalBoundaryError("le calendrier de négociation dépasse le cycle")
    if settlement_dates != (trade_date, settlement_date):
        raise HistoricalBoundaryError("le calendrier de règlement dépasse le cycle")
    calculated = SyntheticSettlementCalendar(settlement_dates).settlement_date(
        trade_date
    )
    if calculated != settlement_date:
        raise HistoricalBoundaryError("la settlement date T+1 n'est pas reproductible")

    sources = _mapping(root["sources"], "sources")
    expected_source_labels = {
        "SPY",
        "BIL",
        "market_audit",
        "distributions",
        "settlement_rule",
        "settlement_calendar",
    }
    if set(sources) != expected_source_labels:
        raise HistoricalBoundaryError("références de sources incomplètes")
    source_ids = {label: _text(sources[label], f"source {label}") for label in sources}

    event_window = _mapping(root["event_window"], "event_window")
    _exact_keys(event_window, {"start", "end", "required_status"}, "event_window")
    event_start = _date(event_window["start"], "event start")
    event_end = _date(event_window["end"], "event end")
    if (
        event_start != decision_date
        or event_end != settlement_date
        or event_window["required_status"] != "CLEAR"
    ):
        raise HistoricalBoundaryError("fenêtre d'événements incohérente")

    return CycleContract(
        digest,
        _text(root["cycle_id"], "cycle_id"),
        _decimal(account["initial_cash"], "cash initial"),
        spy_weight,
        decision_date,
        decision_at,
        trade_date,
        settlement_date,
        trading_dates,
        settlement_dates,
        fees,
        source_ids,
        event_start,
        event_end,
    )


def _load_manifest(repo_root: Path) -> dict[str, Mapping[str, Any]]:
    path = repo_root / "RL" / "data_manifest.json"
    try:
        document = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DocumentaryBlock("manifeste inaccessible ou invalide") from exc
    datasets = document.get("datasets")
    if not isinstance(datasets, list):
        raise DocumentaryBlock("datasets absent du manifeste")
    records = {item.get("id"): item for item in datasets if isinstance(item, dict)}
    if len(records) != len(datasets):
        raise DocumentaryBlock("identifiant de manifeste absent ou dupliqué")
    return records


def _verify_source(
    repo_root: Path, records: Mapping[str, Mapping[str, Any]], identifier: str
) -> VerifiedSource:
    record = records.get(identifier)
    if record is None:
        raise DocumentaryBlock(f"source absente du manifeste: {identifier}")
    expected = record.get("sha256")
    local_file = record.get("local_file")
    if not isinstance(expected, str) or not SHA256_PATTERN.fullmatch(expected):
        raise DocumentaryBlock(f"SHA-256 absent du manifeste: {identifier}")
    if not isinstance(local_file, str):
        raise DocumentaryBlock(f"chemin local absent: {identifier}")
    path = repo_root / local_file
    if not path.is_file():
        raise DocumentaryBlock(f"source locale absente: {identifier}")
    actual = _sha256(path)
    if actual != expected:
        raise HistoricalBoundaryError(f"empreinte source incorrecte: {identifier}")
    expected_bytes = record.get("bytes")
    if not isinstance(expected_bytes, int) or path.stat().st_size != expected_bytes:
        raise HistoricalBoundaryError(f"taille source incorrecte: {identifier}")
    return VerifiedSource(identifier, path, actual, record)


def _load_json_decimal(path: Path) -> Mapping[str, Any]:
    try:
        return _mapping(
            json.loads(path.read_text(encoding="utf-8"), parse_float=Decimal),
            str(path),
        )
    except (OSError, json.JSONDecodeError) as exc:
        raise HistoricalBoundaryError(f"JSON source invalide: {path.name}") from exc


def _timestamp_date(value: Any, timezone_name: str) -> date:
    if type(value) is not int:
        raise HistoricalBoundaryError("timestamp Yahoo invalide")
    return datetime.fromtimestamp(value, tz=ZoneInfo(timezone_name)).date()


def _extract_price(
    document: Mapping[str, Any], symbol: str, price_date: date, field: str
) -> Decimal:
    try:
        result = document["chart"]["result"][0]
        metadata = result["meta"]
        timestamps = result["timestamp"]
        quote = result["indicators"]["quote"][0]
    except (KeyError, IndexError, TypeError) as exc:
        raise HistoricalBoundaryError(f"structure Yahoo invalide pour {symbol}") from exc
    if (
        metadata.get("symbol") != symbol
        or metadata.get("currency") != "USD"
        or metadata.get("exchangeTimezoneName") != "America/New_York"
        or metadata.get("dataGranularity") != "1d"
    ):
        raise HistoricalBoundaryError(f"métadonnées Yahoo invalides pour {symbol}")
    dates = [
        _timestamp_date(item, "America/New_York") for item in timestamps
    ]
    matches = [index for index, item in enumerate(dates) if item == price_date]
    if len(matches) != 1:
        raise DocumentaryBlock(f"observation {symbol} {price_date} absente ou dupliquée")
    index = matches[0]
    try:
        values = {
            name: Decimal(str(quote[name][index]))
            for name in ("open", "high", "low", "close")
        }
    except (KeyError, IndexError, InvalidOperation, TypeError) as exc:
        raise DocumentaryBlock(f"OHLC {symbol} {price_date} invalide") from exc
    if (
        any(not item.is_finite() or item <= 0 for item in values.values())
        or values["low"] > min(values["open"], values["close"])
        or values["high"] < max(values["open"], values["close"])
    ):
        raise DocumentaryBlock(f"OHLC {symbol} {price_date} incohérent")
    return values[field]


def _event_dates(document: Mapping[str, Any]) -> list[date]:
    try:
        result = document["chart"]["result"][0]
    except (KeyError, IndexError, TypeError) as exc:
        raise HistoricalBoundaryError("structure d'événements Yahoo invalide") from exc
    events = result.get("events") or {}
    found: list[date] = []
    for group_name in ("dividends", "splits"):
        group = events.get(group_name) or {}
        if not isinstance(group, dict):
            raise HistoricalBoundaryError("groupe d'événements Yahoo invalide")
        for event in group.values():
            found.append(_timestamp_date(event.get("date"), "America/New_York"))
    return found


def _excel_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%m/%d/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                pass
    raise HistoricalBoundaryError("date de distribution illisible")


def _distribution_events(path: Path, start: date, end: date) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "dividend" not in workbook.sheetnames:
        raise HistoricalBoundaryError("feuille dividend absente")
    rows = workbook["dividend"].iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        raise HistoricalBoundaryError("fichier de distributions vide")
    positions = {str(value): index for index, value in enumerate(headers)}
    required = {"TICKER", "EX-DATE", "RECORD DATE", "PAYABLE DATE"}
    if not required.issubset(positions):
        raise HistoricalBoundaryError("colonnes de distribution absentes")
    found: list[str] = []
    for row in rows:
        ticker = row[positions["TICKER"]]
        if ticker not in SYMBOLS:
            continue
        for field in ("EX-DATE", "RECORD DATE", "PAYABLE DATE"):
            event_date = _excel_date(row[positions[field]])
            if start <= event_date <= end:
                found.append(f"{ticker}:{field}:{event_date.isoformat()}")
    return found


def prepare_historical_cycle(
    config_path: str | Path, *, repo_root: str | Path
) -> PreparedHistoricalCycle:
    contract = load_cycle_contract(config_path)
    root = Path(repo_root)
    records = _load_manifest(root)
    sources = {
        label: _verify_source(root, records, identifier)
        for label, identifier in contract.source_ids.items()
    }
    if sources["settlement_rule"].record.get("status") != "VALIDATED_FOR_USE":
        raise DocumentaryBlock("règle T+1 non validée")
    if sources["settlement_calendar"].record.get("status") != "VALIDATED_FOR_USE":
        raise DocumentaryBlock("preuve DTCC du règlement non validée")
    calendar_record = sources["settlement_calendar"].record
    evidence_scope = " ".join(
        [
            str(calendar_record.get("coverage", "")),
            str(calendar_record.get("economic_meaning", "")),
            " ".join(calendar_record.get("allowed_uses", [])),
        ]
    )
    if "29 mai 2024" not in evidence_scope or "30 mai 2024" not in evidence_scope:
        raise DocumentaryBlock("preuve DTCC non ciblée sur le 29→30 mai 2024")

    documents = {
        symbol: _load_json_decimal(sources[symbol].path) for symbol in SYMBOLS
    }
    audit = _load_json_decimal(sources["market_audit"].path)
    try:
        alert_dates = set(
            audit["cross_checks"]["spy_yahoo_vs_nasdaq_priority_window"]
            ["diagnostic_alerts"]["dates"]
        )
    except (KeyError, TypeError) as exc:
        raise DocumentaryBlock("statuts d'anomalie du marché absents") from exc
    if contract.decision_date.isoformat() in alert_dates or contract.trade_date.isoformat() in alert_dates:
        raise DocumentaryBlock("prix contesté ou signalé pour le cycle")

    event_hits: list[str] = []
    for symbol in SYMBOLS:
        event_hits.extend(
            f"{symbol}:Yahoo:{item.isoformat()}"
            for item in _event_dates(documents[symbol])
            if contract.event_start <= item <= contract.event_end
        )
    event_hits.extend(
        _distribution_events(
            sources["distributions"].path,
            contract.event_start,
            contract.event_end,
        )
    )
    if event_hits:
        raise DocumentaryBlock(
            "événement de distribution ou opération sur titres dans le cycle: "
            + ", ".join(event_hits)
        )

    trace = ["CONFIG_VERIFIED", "SOURCE_HASHES_VERIFIED", "DOCUMENTS_CLEAR"]
    closes = {
        symbol: _extract_price(
            documents[symbol], symbol, contract.decision_date, "close"
        )
        for symbol in SYMBOLS
    }
    trace.append("DECISION_CLOSES_LOADED")
    ledger = SyntheticPortfolioLedger(
        contract.initial_cash,
        strict_temporal=True,
        trading_calendar=SyntheticTradingCalendar(contract.trading_dates),
        settlement_calendar=SyntheticSettlementCalendar(contract.settlement_dates),
    )
    adapter = SyntheticDecisionAdapter(ledger)
    intent = adapter.prepare_intent(
        decision_date=contract.decision_date,
        known_market_wealth=contract.initial_cash,
        inherited_positions={"SPY": 0, "BIL": 0},
        decision_closes={
            symbol: DatedPrice(
                contract.decision_date, closes[symbol], PriceStatus.AVAILABLE
            )
            for symbol in SYMBOLS
        },
        target_weight_spy=contract.target_weight_spy,
    )
    trace.append("INTENT_PREPARED_AND_QUANTITIES_FROZEN")
    return PreparedHistoricalCycle(
        contract,
        sources,
        documents,
        ledger,
        adapter,
        intent,
        closes,
        trace,
    )


def execute_prepared_cycle(prepared: PreparedHistoricalCycle) -> HistoricalCycleResult:
    contract = prepared.contract
    opens = {
        symbol: _extract_price(
            prepared.source_documents[symbol], symbol, contract.trade_date, "open"
        )
        for symbol in SYMBOLS
    }
    prepared.trace.append("EXECUTION_OPENS_LOADED_AFTER_INTENT")
    frozen_quantities = dict(prepared.intent.target_quantities)
    outcome = prepared.adapter.process_intent(
        prepared.intent,
        execution_date=contract.trade_date,
        execution_opens={
            symbol: DatedPrice(
                contract.trade_date, opens[symbol], PriceStatus.AVAILABLE
            )
            for symbol in SYMBOLS
        },
        fees=contract.fees,
    )
    prepared.trace.append("INTENT_PROCESSED_WITHOUT_RESIZING")
    if dict(outcome.intent.target_quantities) != frozen_quantities:
        raise HistoricalBoundaryError("quantités modifiées après l'ouverture")

    source_hashes = {
        label: source.sha256 for label, source in prepared.sources.items()
    }
    if outcome.status is ExecutionStatus.NOT_EVALUABLE:
        return HistoricalCycleResult(
            HistoricalCycleVerdict.NOT_EVALUABLE,
            outcome.reason,
            contract.cycle_id,
            contract.config_sha256,
            tuple(prepared.trace),
            source_hashes,
            prepared.closes,
            opens,
            prepared.intent,
            outcome,
        )
    if outcome.status not in {
        ExecutionStatus.EXECUTED_HYPOTHETICAL,
        ExecutionStatus.ECONOMIC_ORDER_REJECTED,
    }:
        return HistoricalCycleResult(
            HistoricalCycleVerdict.FAILED,
            outcome.reason,
            contract.cycle_id,
            contract.config_sha256,
            tuple(prepared.trace),
            source_hashes,
            prepared.closes,
            opens,
            prepared.intent,
            outcome,
        )
    if outcome.status is ExecutionStatus.ECONOMIC_ORDER_REJECTED:
        prepared.ledger.assert_invariants()
        return HistoricalCycleResult(
            HistoricalCycleVerdict.VALIDATED,
            "rejet économique atomique correctement traité",
            contract.cycle_id,
            contract.config_sha256,
            tuple(prepared.trace),
            source_hashes,
            prepared.closes,
            opens,
            prepared.intent,
            outcome,
            prepared.ledger.accounting_snapshot(),
            prepared.ledger.accounting_snapshot(),
        )

    if outcome.settlement_date != contract.settlement_date:
        raise HistoricalBoundaryError("settlement date différente du calendrier prouvé")
    before_settlement = prepared.ledger.accounting_snapshot()
    wealth_before = prepared.ledger.wealth_market()
    settled_ids = tuple(prepared.ledger.settle_through(contract.settlement_date))
    prepared.trace.append("SETTLEMENT_APPLIED_ON_PROVED_DATE")
    after_settlement = prepared.ledger.accounting_snapshot()
    wealth_after = prepared.ledger.wealth_market()
    prepared.ledger.assert_invariants()
    if wealth_before != wealth_after:
        raise HistoricalBoundaryError("la richesse change pendant le règlement")
    if (
        after_settlement.settled_cash < 0
        or after_settlement.purchase_payables != 0
        or after_settlement.sale_receivables != 0
        or after_settlement.pending_settlement_count != 0
        or after_settlement.fees_recognized != 0
        or dict(after_settlement.economic_positions) != frozen_quantities
        or dict(after_settlement.settled_positions) != frozen_quantities
    ):
        raise HistoricalBoundaryError("invariants finaux du cycle violés")
    prepared.trace.append("FINANCIAL_INVARIANTS_VERIFIED")
    return HistoricalCycleResult(
        HistoricalCycleVerdict.VALIDATED,
        "lot hypothétique et règlement traités correctement",
        contract.cycle_id,
        contract.config_sha256,
        tuple(prepared.trace),
        source_hashes,
        prepared.closes,
        opens,
        prepared.intent,
        outcome,
        before_settlement,
        after_settlement,
        settled_ids,
        wealth_before,
        wealth_after,
    )


def run_single_cycle(
    config_path: str | Path, *, repo_root: str | Path
) -> HistoricalCycleResult:
    try:
        prepared = prepare_historical_cycle(config_path, repo_root=repo_root)
        return execute_prepared_cycle(prepared)
    except DocumentaryBlock as exc:
        return HistoricalCycleResult(
            HistoricalCycleVerdict.NOT_EVALUABLE,
            str(exc),
        )
    except (AdapterError, HistoricalBoundaryError, LedgerError) as exc:
        return HistoricalCycleResult(HistoricalCycleVerdict.FAILED, str(exc))


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    repo_root = Path(__file__).resolve().parents[2]
    config = (
        repo_root
        / "RL"
        / "tests"
        / "fixtures"
        / "historical_spy_bil_single_cycle_v1.json"
    )
    result = run_single_cycle(config, repo_root=repo_root)
    print(json.dumps(result.to_record(), ensure_ascii=False, indent=2, sort_keys=True))
    if result.verdict is HistoricalCycleVerdict.VALIDATED:
        return 0
    if result.verdict is HistoricalCycleVerdict.NOT_EVALUABLE:
        return 2
    return 1


if __name__ == "__main__":
    sys.exit(main())


__all__ = [
    "CycleContract",
    "DocumentaryBlock",
    "HistoricalBoundaryError",
    "HistoricalCycleResult",
    "HistoricalCycleVerdict",
    "PreparedHistoricalCycle",
    "execute_prepared_cycle",
    "load_cycle_contract",
    "prepare_historical_cycle",
    "run_single_cycle",
]
