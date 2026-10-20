#!/usr/bin/env python3
"""Évalue A, B et C1 sur l'unique épisode historique prescrit.

Ce module ne choisit ni dates, ni poids, ni paramètres. Il construit trois achats
initiaux, puis conserve les positions en comptabilisant ex post les distributions
réalisées. Il ne contient aucune boucle de décision ou de rééquilibrage.
"""

from __future__ import annotations

import hashlib
import json
import math
import statistics
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from enum import Enum
from pathlib import Path
from typing import Any, Mapping, Sequence
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from historical_spy_bil_single_cycle import (
    HistoricalBoundaryError,
    VerifiedSource,
    _load_json_decimal,
    _load_manifest,
    _verify_source,
)
from synthetic_decision_adapter import (
    AdapterError,
    DatedPrice,
    ExecutionStatus,
    PriceStatus,
    SyntheticDecisionAdapter,
)
from synthetic_portfolio_ledger import (
    LedgerError,
    SyntheticPortfolioLedger,
    SyntheticSettlementCalendar,
    SyntheticTradingCalendar,
)


SCHEMA_VERSION = "exploratory_buy_hold_spy_bil/v1"
PROVENANCE = "HISTORICAL_AUDITED_PROXY"
SYMBOLS = ("SPY", "BIL")
POLICY_IDS = ("A", "B", "C1")
ZERO = Decimal("0")


class EvaluationVerdict(str, Enum):
    VALIDATED = "EXPLORATORY_EVALUATION_VALIDATED"
    PARTIAL = "PARTIALLY_EVALUABLE"
    NOT_EVALUABLE = "NOT_EVALUABLE"
    FAILED = "FAILED_VALIDATION"


class DocumentaryBlock(HistoricalBoundaryError):
    """Entrée commune indispensable absente ou contestée."""


@dataclass(frozen=True)
class PolicyDefinition:
    policy_id: str
    label: str
    target_weight_spy: Decimal


@dataclass(frozen=True)
class EpisodeContract:
    episode_id: str
    config_sha256: str
    decision_date: date
    decision_at: datetime
    execution_date: date
    settlement_date: date
    anniversary_date: date
    initial_cash: Decimal
    policies: tuple[PolicyDefinition, ...]
    source_ids: Mapping[str, str]
    volatility_factor: int
    annualized_return_day_count: int


@dataclass(frozen=True)
class DistributionEvent:
    event_id: str
    symbol: str
    ex_date: date
    record_date: date
    payment_date: date
    amount_per_share: Decimal


@dataclass(frozen=True)
class EpisodeData:
    terminal_date: date
    common_sessions: tuple[date, ...]
    decision_closes: Mapping[str, Decimal]
    execution_opens: Mapping[str, Decimal]
    closing_prices: Mapping[date, Mapping[str, Decimal]]
    distributions: tuple[DistributionEvent, ...]
    sources: Mapping[str, VerifiedSource]


@dataclass(frozen=True)
class ValuationRow:
    label: str
    valuation_date: date
    wealth: Decimal
    net_cash: Decimal
    distribution_receivable: Decimal
    spy_weight: Decimal
    bil_weight: Decimal

    def to_record(self) -> dict[str, Any]:
        return {
            "label": self.label,
            "date": self.valuation_date.isoformat(),
            "wealth": decimal_text(self.wealth),
            "net_cash": decimal_text(self.net_cash),
            "distribution_receivable": decimal_text(
                self.distribution_receivable
            ),
            "spy_weight": decimal_text(self.spy_weight),
            "bil_weight": decimal_text(self.bil_weight),
        }


@dataclass(frozen=True)
class PolicyResult:
    policy_id: str
    label: str
    status: str
    target_weight_spy: Decimal
    quantities: Mapping[str, int]
    initial_orders: int
    initial_lots: int
    initial_obtained_weights: Mapping[str, Decimal]
    terminal_weights: Mapping[str, Decimal]
    spy_weight_min: Decimal
    spy_weight_max: Decimal
    mean_spy_exposure: Decimal
    initial_wealth: Decimal
    terminal_wealth: Decimal
    total_return: float
    annualized_return: float
    annualized_volatility: float
    worst_daily_return: float
    maximum_drawdown: float
    maximum_drawdown_duration_sessions: int
    terminal_drawdown_open: bool
    mean_net_cash: Decimal
    terminal_cash: Decimal
    distribution_contributions: Mapping[str, Decimal]
    distribution_event_count: int
    positive_distribution_entitlements: int
    economic_rejections: int
    settled_ids: tuple[int, ...]
    trajectory: tuple[ValuationRow, ...]

    def to_record(self, *, include_trajectory: bool = False) -> dict[str, Any]:
        record: dict[str, Any] = {
            "policy_id": self.policy_id,
            "label": self.label,
            "status": self.status,
            "target_weight_spy": decimal_text(self.target_weight_spy),
            "quantities": dict(self.quantities),
            "initial_orders": self.initial_orders,
            "initial_lots": self.initial_lots,
            "initial_obtained_weights": {
                key: decimal_text(value)
                for key, value in self.initial_obtained_weights.items()
            },
            "terminal_weights": {
                key: decimal_text(value)
                for key, value in self.terminal_weights.items()
            },
            "spy_weight_min": decimal_text(self.spy_weight_min),
            "spy_weight_max": decimal_text(self.spy_weight_max),
            "mean_spy_exposure": decimal_text(self.mean_spy_exposure),
            "initial_wealth": decimal_text(self.initial_wealth),
            "terminal_wealth": decimal_text(self.terminal_wealth),
            "total_return": float_text(self.total_return),
            "annualized_return": float_text(self.annualized_return),
            "annualized_volatility": float_text(self.annualized_volatility),
            "worst_daily_return": float_text(self.worst_daily_return),
            "maximum_drawdown": float_text(self.maximum_drawdown),
            "maximum_drawdown_duration_sessions": (
                self.maximum_drawdown_duration_sessions
            ),
            "terminal_drawdown_open": self.terminal_drawdown_open,
            "mean_net_cash": decimal_text(self.mean_net_cash),
            "terminal_cash": decimal_text(self.terminal_cash),
            "distribution_contributions": {
                key: decimal_text(value)
                for key, value in self.distribution_contributions.items()
            },
            "distribution_event_count": self.distribution_event_count,
            "positive_distribution_entitlements": (
                self.positive_distribution_entitlements
            ),
            "economic_rejections": self.economic_rejections,
            "settled_ids": list(self.settled_ids),
            "valuation_count_including_entry": len(self.trajectory),
        }
        if include_trajectory:
            record["trajectory"] = [row.to_record() for row in self.trajectory]
        return record


@dataclass(frozen=True)
class EvaluationResult:
    verdict: EvaluationVerdict
    reason: str
    episode_id: str | None = None
    config_sha256: str | None = None
    terminal_date: date | None = None
    source_hashes: Mapping[str, str] | None = None
    market_summary: Mapping[str, Any] | None = None
    policies: tuple[PolicyResult, ...] = ()
    trace: tuple[str, ...] = ()

    def to_record(self, *, include_trajectories: bool = False) -> dict[str, Any]:
        return {
            "verdict": self.verdict.value,
            "reason": self.reason,
            "episode_id": self.episode_id,
            "config_sha256": self.config_sha256,
            "terminal_date": (
                None if self.terminal_date is None else self.terminal_date.isoformat()
            ),
            "source_hashes": dict(self.source_hashes or {}),
            "market_summary": dict(self.market_summary or {}),
            "trace": list(self.trace),
            "policies": [
                item.to_record(include_trajectory=include_trajectories)
                for item in self.policies
            ],
        }


def decimal_text(value: Decimal) -> str:
    return format(value, "f")


def float_text(value: float) -> str:
    return format(value, ".12f")


def _mapping(value: Any, label: str) -> Mapping[str, Any]:
    if not isinstance(value, dict):
        raise HistoricalBoundaryError(f"{label} doit être un objet")
    return value


def _exact_keys(value: Mapping[str, Any], expected: set[str], label: str) -> None:
    if set(value) != expected:
        raise HistoricalBoundaryError(f"champs {label} inattendus ou manquants")


def _date(value: Any, label: str) -> date:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise HistoricalBoundaryError(f"{label} invalide") from exc


def _decimal(value: Any, label: str) -> Decimal:
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise HistoricalBoundaryError(f"{label} invalide") from exc
    if not result.is_finite():
        raise HistoricalBoundaryError(f"{label} doit être fini")
    return result


def load_contract(config_path: str | Path) -> EpisodeContract:
    path = Path(config_path)
    raw = path.read_bytes()
    sidecar = path.with_suffix(path.suffix + ".sha256")
    try:
        expected_hash = sidecar.read_text(encoding="ascii").split()[0]
        root = _mapping(json.loads(raw.decode("utf-8")), "racine")
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, IndexError) as exc:
        raise HistoricalBoundaryError("configuration ou sidecar illisible") from exc
    actual_hash = hashlib.sha256(raw).hexdigest()
    if actual_hash != expected_hash:
        raise HistoricalBoundaryError("empreinte de configuration incorrecte")
    _exact_keys(
        root,
        {
            "schema_version",
            "episode_id",
            "provenance",
            "scope",
            "dates",
            "account",
            "policies",
            "prices",
            "distributions",
            "metrics",
            "coverage",
            "sources",
        },
        "racine",
    )
    if root["schema_version"] != SCHEMA_VERSION or root["provenance"] != PROVENANCE:
        raise HistoricalBoundaryError("schéma ou provenance historique invalide")
    if root["scope"] != "ONE_PRESPECIFIED_EPISODE_A_B_C1":
        raise HistoricalBoundaryError("le périmètre doit rester l'épisode A/B/C1")

    dates = _mapping(root["dates"], "dates")
    _exact_keys(
        dates,
        {
            "decision_price_date",
            "decision_at",
            "execution_date",
            "initial_settlement_date",
            "anniversary_date",
            "terminal_rule",
        },
        "dates",
    )
    decision_date = _date(dates["decision_price_date"], "date de close")
    execution_date = _date(dates["execution_date"], "date d'exécution")
    settlement_date = _date(dates["initial_settlement_date"], "date de règlement")
    anniversary_date = _date(dates["anniversary_date"], "date anniversaire")
    try:
        decision_at = datetime.fromisoformat(dates["decision_at"])
        local = decision_at.astimezone(ZoneInfo("America/New_York"))
    except (TypeError, ValueError) as exc:
        raise HistoricalBoundaryError("instant de décision invalide") from exc
    if decision_at.tzinfo is None or local.strftime("%H:%M:%S") != "08:00:00":
        raise HistoricalBoundaryError("la décision doit être zonée à 08:00 ET")
    if not (
        decision_date == date(2024, 5, 28)
        and execution_date == date(2024, 5, 29)
        and settlement_date == date(2024, 5, 30)
        and anniversary_date == date(2025, 5, 29)
        and dates["terminal_rule"]
        == "LAST_COMMON_CLOSE_ON_OR_BEFORE_ANNIVERSARY"
    ):
        raise HistoricalBoundaryError("dates préspécifiées modifiées")

    account = _mapping(root["account"], "account")
    if account != {
        "currency": "USD",
        "initial_cash": "100000",
        "strict_temporal": True,
        "whole_shares": True,
        "funding": "H_FIN_NET",
        "external_fees": "0",
        "residual_cash_interest": "0",
        "automatic_distribution_reinvestment": False,
        "forced_terminal_liquidation": False,
    }:
        raise HistoricalBoundaryError("conventions comptables modifiées")

    policies_raw = root["policies"]
    if not isinstance(policies_raw, list) or len(policies_raw) != 3:
        raise HistoricalBoundaryError("trois politiques sont requises")
    policies = tuple(
        PolicyDefinition(
            str(item["id"]),
            str(item["label"]),
            _decimal(item["target_weight_spy"], "poids SPY"),
        )
        for item in policies_raw
    )
    if tuple(item.policy_id for item in policies) != POLICY_IDS or tuple(
        item.target_weight_spy for item in policies
    ) != (Decimal("1"), Decimal("0"), Decimal("0.5")):
        raise HistoricalBoundaryError("politiques ou poids modifiés")

    prices = _mapping(root["prices"], "prices")
    if prices != {
        "decision_field": "close",
        "execution_field": "open",
        "valuation_field": "close",
        "status": "MARKET_OBSERVABLE_PROXY",
        "adjusted_close_prohibited": True,
    }:
        raise HistoricalBoundaryError("convention de prix modifiée")
    distributions = _mapping(root["distributions"], "distributions")
    if distributions != {
        "mode": "REALIZED_EX_POST_ACCOUNTING_ONLY",
        "entitlement_at_ex_date": True,
        "cash_at_payment_date": True,
        "no_decision_after_initial_purchase": True,
    }:
        raise HistoricalBoundaryError("convention de distribution modifiée")
    coverage = _mapping(root["coverage"], "coverage")
    if coverage != {
        "comparison": "PAIRED",
        "block_all_policies_on_required_common_source_failure": True,
    }:
        raise HistoricalBoundaryError("couverture appariée modifiée")
    metrics = _mapping(root["metrics"], "metrics")
    if metrics != {
        "daily_return_first_interval": "EXECUTION_OPEN_TO_FIRST_CLOSE",
        "daily_return_later_intervals": "COMMON_CLOSE_TO_COMMON_CLOSE",
        "volatility_annualization_factor": 252,
        "annualized_return_day_count": 365,
        "drawdown_basis": "ECONOMIC_WEALTH_WITH_ENTRY_OPEN_BASE",
    }:
        raise HistoricalBoundaryError("conventions de mesure modifiées")
    sources = _mapping(root["sources"], "sources")
    if set(sources) != {
        "SPY",
        "BIL",
        "market_audit",
        "distributions",
        "settlement_rule",
        "settlement_calendar",
    } or any(not isinstance(value, str) for value in sources.values()):
        raise HistoricalBoundaryError("références de sources invalides")
    return EpisodeContract(
        str(root["episode_id"]),
        actual_hash,
        decision_date,
        decision_at,
        execution_date,
        settlement_date,
        anniversary_date,
        Decimal("100000"),
        policies,
        dict(sources),
        252,
        365,
    )


def _timestamp_date(timestamp: Any) -> date:
    if isinstance(timestamp, bool) or not isinstance(timestamp, int):
        raise DocumentaryBlock("timestamp Yahoo invalide")
    return datetime.fromtimestamp(
        timestamp, ZoneInfo("America/New_York")
    ).date()


def _yahoo_prices(document: Mapping[str, Any], symbol: str) -> dict[date, dict[str, Decimal]]:
    try:
        result = document["chart"]["result"][0]
        meta = result["meta"]
        timestamps = result["timestamp"]
        quote = result["indicators"]["quote"][0]
    except (KeyError, IndexError, TypeError) as exc:
        raise DocumentaryBlock(f"structure Yahoo {symbol} invalide") from exc
    if (
        meta.get("symbol") != symbol
        or meta.get("currency") != "USD"
        or meta.get("exchangeTimezoneName") != "America/New_York"
        or meta.get("dataGranularity") != "1d"
    ):
        raise DocumentaryBlock(f"métadonnées Yahoo {symbol} incompatibles")
    if len(timestamps) != len(quote["open"]) or len(timestamps) != len(quote["close"]):
        raise DocumentaryBlock(f"longueurs Yahoo {symbol} incohérentes")
    found: dict[date, dict[str, Decimal]] = {}
    for index, timestamp in enumerate(timestamps):
        current = _timestamp_date(timestamp)
        try:
            values = {
                field: Decimal(str(quote[field][index]))
                for field in ("open", "high", "low", "close")
            }
        except (KeyError, IndexError, InvalidOperation, TypeError) as exc:
            raise DocumentaryBlock(f"OHLC {symbol} invalide") from exc
        if any(not value.is_finite() or value <= 0 for value in values.values()):
            raise DocumentaryBlock(f"OHLC {symbol} non positif")
        if current in found:
            raise DocumentaryBlock(f"date Yahoo {symbol} dupliquée")
        found[current] = values
    return found


def _yahoo_events(document: Mapping[str, Any]) -> tuple[dict[date, Decimal], list[date]]:
    try:
        events = document["chart"]["result"][0].get("events") or {}
        dividends = {
            _timestamp_date(item["date"]): Decimal(str(item["amount"]))
            for item in (events.get("dividends") or {}).values()
        }
        splits = [
            _timestamp_date(item["date"])
            for item in (events.get("splits") or {}).values()
        ]
    except (KeyError, InvalidOperation, TypeError) as exc:
        raise DocumentaryBlock("événements Yahoo invalides") from exc
    return dividends, splits


def _excel_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%m/%d/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                pass
    raise DocumentaryBlock("date State Street illisible")


def _cell_decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return ZERO
    return _decimal(value, "montant de distribution")


def _load_distributions(
    path: Path, start: date, end: date
) -> tuple[DistributionEvent, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "dividend" not in workbook.sheetnames:
            raise DocumentaryBlock("feuille de distributions absente")
        rows = workbook["dividend"].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise DocumentaryBlock("fichier de distributions vide")
        columns = {str(value): index for index, value in enumerate(headers)}
        required = {
            "TICKER",
            "EX-DATE",
            "RECORD DATE",
            "PAYABLE DATE",
            "DIVIDEND ($)",
            "SHORT TERM CAPITAL GAIN ($)",
            "LONG TERM CAPITAL GAIN ($)",
        }
        if not required.issubset(columns):
            raise DocumentaryBlock("colonnes State Street absentes")
        events: list[DistributionEvent] = []
        for row in rows:
            symbol = row[columns["TICKER"]]
            if symbol not in SYMBOLS:
                continue
            ex_date = _excel_date(row[columns["EX-DATE"]])
            if not start <= ex_date <= end:
                continue
            record_date = _excel_date(row[columns["RECORD DATE"]])
            payment_date = _excel_date(row[columns["PAYABLE DATE"]])
            amount = sum(
                (
                    _cell_decimal(row[columns[field]])
                    for field in (
                        "DIVIDEND ($)",
                        "SHORT TERM CAPITAL GAIN ($)",
                        "LONG TERM CAPITAL GAIN ($)",
                    )
                ),
                ZERO,
            )
            if amount <= 0 or record_date < ex_date or payment_date < ex_date:
                raise DocumentaryBlock("distribution State Street incohérente")
            events.append(
                DistributionEvent(
                    f"SSGA-{symbol}-{ex_date.isoformat()}",
                    symbol,
                    ex_date,
                    record_date,
                    payment_date,
                    amount,
                )
            )
    finally:
        workbook.close()
    events.sort(key=lambda item: (item.ex_date, item.symbol))
    identifiers = [item.event_id for item in events]
    if len(set(identifiers)) != len(identifiers):
        raise DocumentaryBlock("distribution State Street dupliquée")
    return tuple(events)


def prepare_episode_data(
    contract: EpisodeContract, *, repo_root: str | Path
) -> EpisodeData:
    root = Path(repo_root)
    records = _load_manifest(root)
    sources = {
        label: _verify_source(root, records, identifier)
        for label, identifier in contract.source_ids.items()
    }
    required_statuses = {
        "SPY": "AUDITED",
        "BIL": "AUDITED",
        "market_audit": "AUDITED",
        "distributions": "VALIDATED_FOR_USE",
    }
    for label, expected in required_statuses.items():
        if sources[label].record.get("status") != expected:
            raise DocumentaryBlock(f"statut de source {label} incompatible")
    if sources["settlement_rule"].record.get("status") != "VALIDATED_FOR_USE":
        raise DocumentaryBlock("règle T+1 non validée")
    if sources["settlement_calendar"].record.get("status") != "VALIDATED_FOR_USE":
        raise DocumentaryBlock("règlement initial non documenté")
    documents = {
        symbol: _load_json_decimal(sources[symbol].path) for symbol in SYMBOLS
    }
    prices = {
        symbol: _yahoo_prices(documents[symbol], symbol) for symbol in SYMBOLS
    }
    common = tuple(
        sorted(
            current
            for current in set(prices["SPY"]).intersection(prices["BIL"])
            if contract.decision_date <= current <= contract.anniversary_date
        )
    )
    if not common or common[0] != contract.decision_date:
        raise DocumentaryBlock("close de décision absent du calendrier commun")
    if contract.execution_date not in common:
        raise DocumentaryBlock("date d'exécution absente du calendrier commun")
    terminal = max(current for current in common if current <= contract.anniversary_date)
    valuation_sessions = tuple(
        current for current in common if contract.execution_date <= current <= terminal
    )
    if terminal != contract.anniversary_date:
        raise DocumentaryBlock("la date terminale commune ne peut pas être établie")

    audit = _load_json_decimal(sources["market_audit"].path)
    try:
        alert_dates = {
            date.fromisoformat(value)
            for value in audit["cross_checks"]
            ["spy_yahoo_vs_nasdaq_priority_window"]
            ["diagnostic_alerts"]["dates"]
        }
    except (KeyError, TypeError, ValueError) as exc:
        raise DocumentaryBlock("audit d'anomalies invalide") from exc
    hits = sorted(set(valuation_sessions).intersection(alert_dates))
    if hits:
        raise DocumentaryBlock(
            "prix signalé dans l'épisode: " + ", ".join(item.isoformat() for item in hits)
        )

    distributions = _load_distributions(
        sources["distributions"].path, contract.execution_date, terminal
    )
    for symbol in SYMBOLS:
        yahoo_dividends, yahoo_splits = _yahoo_events(documents[symbol])
        split_hits = [
            item
            for item in yahoo_splits
            if contract.execution_date <= item <= terminal
        ]
        if split_hits:
            raise DocumentaryBlock(f"opération sur titres {symbol} dans l'épisode")
        official = {item.ex_date: item for item in distributions if item.symbol == symbol}
        yahoo_window = {
            event_date: amount
            for event_date, amount in yahoo_dividends.items()
            if contract.execution_date <= event_date <= terminal
        }
        if set(official) != set(yahoo_window):
            raise DocumentaryBlock(f"dates de distributions {symbol} contradictoires")
        for event_date, event in official.items():
            if event.amount_per_share.quantize(Decimal("0.001")) != yahoo_window[
                event_date
            ].quantize(Decimal("0.001")):
                raise DocumentaryBlock(f"montant de distribution {symbol} contradictoire")
            if event.payment_date > terminal:
                raise DocumentaryBlock("distribution terminale encore impayée non prévue")

    return EpisodeData(
        terminal,
        valuation_sessions,
        {symbol: prices[symbol][contract.decision_date]["close"] for symbol in SYMBOLS},
        {symbol: prices[symbol][contract.execution_date]["open"] for symbol in SYMBOLS},
        {
            current: {symbol: prices[symbol][current]["close"] for symbol in SYMBOLS}
            for current in valuation_sessions
        },
        distributions,
        sources,
    )


def _economic_wealth(
    ledger: SyntheticPortfolioLedger, prices: Mapping[str, Decimal]
) -> Decimal:
    holdings = sum(
        (
            prices[symbol] * ledger.economic_positions.get(symbol, 0)
            for symbol in SYMBOLS
        ),
        ZERO,
    )
    return (
        holdings
        + ledger.settled_cash
        + ledger.distribution_receivables
        + ledger.sale_receivables
        - ledger.purchase_payables
    )


def _valuation_row(
    label: str,
    valuation_date: date,
    ledger: SyntheticPortfolioLedger,
    prices: Mapping[str, Decimal],
) -> ValuationRow:
    wealth = _economic_wealth(ledger, prices)
    if wealth <= 0:
        raise HistoricalBoundaryError("richesse non positive")
    net_cash = ledger.settled_cash + ledger.sale_receivables - ledger.purchase_payables
    return ValuationRow(
        label,
        valuation_date,
        wealth,
        net_cash,
        ledger.distribution_receivables,
        prices["SPY"] * ledger.economic_positions.get("SPY", 0) / wealth,
        prices["BIL"] * ledger.economic_positions.get("BIL", 0) / wealth,
    )


def calculate_path_metrics(
    trajectory: Sequence[ValuationRow],
    *,
    volatility_factor: int,
    elapsed_days: int,
    annualized_return_day_count: int,
) -> dict[str, Any]:
    if len(trajectory) < 2 or elapsed_days <= 0:
        raise HistoricalBoundaryError("trajectoire insuffisante")
    wealth = [float(item.wealth) for item in trajectory]
    returns = [wealth[index] / wealth[index - 1] - 1 for index in range(1, len(wealth))]
    total_return = wealth[-1] / wealth[0] - 1
    annualized_return = (wealth[-1] / wealth[0]) ** (
        annualized_return_day_count / elapsed_days
    ) - 1
    annualized_volatility = statistics.stdev(returns) * math.sqrt(volatility_factor)
    peak = wealth[0]
    maximum_drawdown = 0.0
    current_duration = 0
    maximum_duration = 0
    for value in wealth[1:]:
        if value >= peak:
            peak = value
            current_duration = 0
        else:
            current_duration += 1
            maximum_duration = max(maximum_duration, current_duration)
            maximum_drawdown = min(maximum_drawdown, value / peak - 1)
    return {
        "total_return": total_return,
        "annualized_return": annualized_return,
        "annualized_volatility": annualized_volatility,
        "worst_daily_return": min(returns),
        "maximum_drawdown": maximum_drawdown,
        "maximum_drawdown_duration_sessions": maximum_duration,
        "terminal_drawdown_open": current_duration > 0,
    }


def evaluate_policy(
    contract: EpisodeContract,
    data: EpisodeData,
    policy: PolicyDefinition,
) -> PolicyResult:
    trading_dates = (contract.decision_date, *data.common_sessions)
    ledger = SyntheticPortfolioLedger(
        contract.initial_cash,
        strict_temporal=True,
        trading_calendar=SyntheticTradingCalendar(trading_dates),
        settlement_calendar=SyntheticSettlementCalendar(
            [contract.execution_date, contract.settlement_date]
        ),
    )
    adapter = SyntheticDecisionAdapter(ledger)
    intent = adapter.prepare_intent(
        decision_date=contract.decision_date,
        known_market_wealth=contract.initial_cash,
        inherited_positions={"SPY": 0, "BIL": 0},
        decision_closes={
            symbol: DatedPrice(
                contract.decision_date,
                data.decision_closes[symbol],
                PriceStatus.AVAILABLE,
            )
            for symbol in SYMBOLS
        },
        target_weight_spy=policy.target_weight_spy,
    )
    quantities = dict(intent.target_quantities)
    outcome = adapter.process_intent(
        intent,
        execution_date=contract.execution_date,
        execution_opens={
            symbol: DatedPrice(
                contract.execution_date,
                data.execution_opens[symbol],
                PriceStatus.AVAILABLE,
            )
            for symbol in SYMBOLS
        },
        fees={symbol: ZERO for symbol in SYMBOLS},
    )
    if outcome.status is not ExecutionStatus.EXECUTED_HYPOTHETICAL:
        raise HistoricalBoundaryError(
            f"{policy.policy_id}: achat initial non exécuté ({outcome.status.value})"
        )
    if dict(intent.target_quantities) != quantities:
        raise HistoricalBoundaryError("quantités modifiées après l'open")

    trajectory = [
        _valuation_row(
            "EXECUTION_OPEN_BASE",
            contract.execution_date,
            ledger,
            data.execution_opens,
        )
    ]
    initial_weights = {
        "SPY": trajectory[0].spy_weight,
        "BIL": trajectory[0].bil_weight,
        "CASH": trajectory[0].net_cash / trajectory[0].wealth,
    }
    ex_events: dict[date, list[DistributionEvent]] = {}
    payment_dates: set[date] = set()
    for event in data.distributions:
        ex_events.setdefault(event.ex_date, []).append(event)
        payment_dates.add(event.payment_date)
    contributions = {symbol: ZERO for symbol in SYMBOLS}
    settled_ids: tuple[int, ...] = ()

    for current in data.common_sessions:
        if current == contract.execution_date:
            trajectory.append(
                _valuation_row(
                    "COMMON_CLOSE", current, ledger, data.closing_prices[current]
                )
            )
            continue
        for event in ex_events.get(current, []):
            ledger.accrue_realized_distribution(
                event.symbol,
                event_id=event.event_id,
                ex_date=event.ex_date,
                payment_date=event.payment_date,
                amount_per_share=event.amount_per_share,
            )
            contributions[event.symbol] += (
                event.amount_per_share * quantities[event.symbol]
            )
        if current == contract.settlement_date:
            settled_ids = tuple(ledger.settle_through(current))
        if current in payment_dates:
            ledger.pay_distributions_through(current)
        for symbol in SYMBOLS:
            ledger.set_prices(
                symbol,
                market_price=data.closing_prices[current][symbol],
                price_date=current,
            )
        ledger.assert_invariants()
        if dict(ledger.economic_positions) != {
            symbol: quantity for symbol, quantity in quantities.items() if quantity
        }:
            raise HistoricalBoundaryError("position modifiée après l'achat initial")
        trajectory.append(
            _valuation_row(
                "COMMON_CLOSE", current, ledger, data.closing_prices[current]
            )
        )

    if ledger.pending_settlements or ledger.distribution_receivables != 0:
        raise HistoricalBoundaryError("flux non dénoué au terminal")
    if ledger.fees_recognized != 0 or ledger.settled_cash < 0:
        raise HistoricalBoundaryError("frais ou cash terminal incohérent")
    trade_entries = [entry for entry in ledger.journal if entry.kind == "TRADE"]
    if len(trade_entries) != len(intent.orders):
        raise HistoricalBoundaryError("ordre supplémentaire détecté")
    metrics = calculate_path_metrics(
        trajectory,
        volatility_factor=contract.volatility_factor,
        elapsed_days=(data.terminal_date - contract.execution_date).days,
        annualized_return_day_count=contract.annualized_return_day_count,
    )
    closes = trajectory[1:]
    mean_spy = sum((item.spy_weight for item in closes), ZERO) / len(closes)
    mean_cash = sum((item.net_cash for item in trajectory), ZERO) / len(trajectory)
    terminal = trajectory[-1]
    return PolicyResult(
        policy.policy_id,
        policy.label,
        outcome.status.value,
        policy.target_weight_spy,
        quantities,
        len(intent.orders),
        1,
        initial_weights,
        {
            "SPY": terminal.spy_weight,
            "BIL": terminal.bil_weight,
            "CASH": terminal.net_cash / terminal.wealth,
        },
        min(item.spy_weight for item in closes),
        max(item.spy_weight for item in closes),
        mean_spy,
        trajectory[0].wealth,
        terminal.wealth,
        metrics["total_return"],
        metrics["annualized_return"],
        metrics["annualized_volatility"],
        metrics["worst_daily_return"],
        metrics["maximum_drawdown"],
        metrics["maximum_drawdown_duration_sessions"],
        metrics["terminal_drawdown_open"],
        mean_cash,
        ledger.settled_cash,
        contributions,
        len(data.distributions),
        sum(quantities[event.symbol] > 0 for event in data.distributions),
        len(ledger.rejections),
        settled_ids,
        tuple(trajectory),
    )


def run_evaluation(
    config_path: str | Path, *, repo_root: str | Path
) -> EvaluationResult:
    trace: list[str] = []
    try:
        contract = load_contract(config_path)
        trace.append("CONFIG_VERIFIED")
        data = prepare_episode_data(contract, repo_root=repo_root)
        trace.extend(
            [
                "SOURCE_HASHES_VERIFIED",
                "PAIRED_MARKET_COVERAGE_VERIFIED",
                "REALIZED_DISTRIBUTIONS_RECONCILED",
                "TERMINAL_COMMON_CLOSE_FIXED",
            ]
        )
        policies = tuple(
            evaluate_policy(contract, data, policy) for policy in contract.policies
        )
        trace.extend(
            [
                "INITIAL_INTENTS_FROZEN_BEFORE_OPENS",
                "INITIAL_LOTS_SETTLED_T_PLUS_ONE",
                "NO_POST_INITIAL_ORDERS",
                "ACCOUNTING_AND_PATH_METRICS_VERIFIED",
            ]
        )
        return EvaluationResult(
            EvaluationVerdict.VALIDATED,
            "épisode apparié A/B/C1 entièrement évalué",
            contract.episode_id,
            contract.config_sha256,
            data.terminal_date,
            {label: source.sha256 for label, source in data.sources.items()},
            {
                "common_close_count": len(data.common_sessions),
                "decision_closes": {
                    symbol: decimal_text(data.decision_closes[symbol])
                    for symbol in SYMBOLS
                },
                "execution_opens": {
                    symbol: decimal_text(data.execution_opens[symbol])
                    for symbol in SYMBOLS
                },
                "terminal_closes": {
                    symbol: decimal_text(data.closing_prices[data.terminal_date][symbol])
                    for symbol in SYMBOLS
                },
                "distribution_events": {
                    symbol: sum(
                        event.symbol == symbol for event in data.distributions
                    )
                    for symbol in SYMBOLS
                },
            },
            policies,
            tuple(trace),
        )
    except DocumentaryBlock as exc:
        return EvaluationResult(EvaluationVerdict.NOT_EVALUABLE, str(exc), trace=tuple(trace))
    except (AdapterError, HistoricalBoundaryError, LedgerError) as exc:
        return EvaluationResult(EvaluationVerdict.FAILED, str(exc), trace=tuple(trace))


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    root = Path(__file__).resolve().parents[2]
    config = (
        root
        / "RL"
        / "tests"
        / "fixtures"
        / "exploratory_buy_hold_spy_bil_v1.json"
    )
    result = run_evaluation(config, repo_root=root)
    print(json.dumps(result.to_record(), ensure_ascii=False, indent=2, sort_keys=True))
    return 0 if result.verdict is EvaluationVerdict.VALIDATED else 1


if __name__ == "__main__":
    raise SystemExit(main())
