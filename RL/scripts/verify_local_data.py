#!/usr/bin/env python3
"""Vérifie hors ligne les fichiers bruts déclarés dans le registre RL.

Ce script ne télécharge rien, ne transforme rien et n'impute aucune valeur.
Il contrôle l'intégrité, la lisibilité et quelques invariants propres aux formats.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import sys
import warnings
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


warnings.filterwarnings("ignore", message="Cannot parse header or footer.*", module="openpyxl")


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MANIFEST = ROOT / "RL" / "data_manifest.json"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def iso_date_range(values: pd.Series) -> tuple[str | None, str | None, int]:
    dates = pd.to_datetime(values, errors="coerce")
    valid = dates.dropna()
    if valid.empty:
        return None, None, int(dates.isna().sum())
    return valid.min().date().isoformat(), valid.max().date().isoformat(), int(dates.isna().sum())


def finite_summary(frame: pd.DataFrame) -> dict[str, int]:
    numeric = frame.apply(pd.to_numeric, errors="coerce")
    array = numeric.to_numpy(dtype=float)
    return {
        "missing_numeric": int(pd.isna(numeric).sum().sum()),
        "infinite": int((~pd.isna(array) & ~pd.DataFrame(array).map(math.isfinite).to_numpy()).sum()),
    }


def audit_ff_daily_zip(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        members = archive.namelist()
        if len(members) != 1:
            raise ValueError(f"archive Fama–French inattendue: {members}")
        text = archive.read(members[0]).decode("utf-8-sig")
    lines = text.splitlines()
    header = next(i for i, line in enumerate(lines) if line.strip().startswith(",Mkt-RF"))
    selected = [lines[header]]
    for line in lines[header + 1 :]:
        first = line.split(",", 1)[0].strip()
        if len(first) == 8 and first.isdigit():
            selected.append(line)
        elif len(selected) > 1:
            break
    frame = pd.read_csv(io.StringIO("\n".join(selected))).rename(columns=lambda c: "date" if str(c).startswith("Unnamed") else c)
    dates = pd.to_datetime(frame["date"].astype(str), format="%Y%m%d", errors="raise")
    numeric = frame.drop(columns="date").apply(pd.to_numeric, errors="coerce")
    return {
        "zip_members": members,
        "zip_crc_error": bad_member,
        "rows": len(frame),
        "columns": list(frame.columns),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "missing_values": int(numeric.isna().sum().sum()),
        "infinite_values": int((~pd.DataFrame(numeric).map(math.isfinite) & numeric.notna()).sum().sum()),
    }


def audit_nasdaq_json(path: Path, params: dict[str, Any]) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    data = payload.get("data") or {}
    rows = (data.get("tradesTable") or {}).get("rows") or []
    frame = pd.DataFrame(rows)
    dates = pd.to_datetime(frame["date"], format="%m/%d/%Y", errors="raise")
    numeric = pd.DataFrame(index=frame.index)
    for column in ["open", "high", "low", "close", "volume"]:
        numeric[column] = pd.to_numeric(
            frame[column].astype(str).str.replace("$", "", regex=False).str.replace(",", "", regex=False),
            errors="coerce",
        )
    ohlc = numeric[["open", "close"]].min(axis=1).ge(numeric["low"]) & numeric[["open", "close"]].max(axis=1).le(numeric["high"])
    return {
        "ticker": params["ticker"],
        "rows": len(frame),
        "provider_total_records": int(data.get("totalRecords", -1)),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "weekend_dates": int((dates.dt.dayofweek >= 5).sum()),
        "missing_ohlcv": int(numeric.isna().sum().sum()),
        "invalid_ohlc_rows": int((~ohlc).sum()),
        "negative_volume_rows": int((numeric["volume"] < 0).sum()),
    }


def audit_fred_rates(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(path, na_values=["."])
    date_column = frame.columns[0]
    dates = pd.to_datetime(frame[date_column], errors="raise")
    values = frame.drop(columns=date_column).apply(pd.to_numeric, errors="coerce")
    coverage = {}
    for column in values:
        mask = values[column].notna()
        coverage[column] = {
            "nonmissing": int(mask.sum()),
            "start": dates[mask].min().date().isoformat(),
            "end": dates[mask].max().date().isoformat(),
            "missing": int((~mask).sum()),
        }
    return {
        "rows": len(frame),
        "duplicate_dates": int(dates.duplicated().sum()),
        "columns": list(values.columns),
        "coverage": coverage,
        "infinite_values": int((~pd.DataFrame(values).map(math.isfinite) & values.notna()).sum().sum()),
    }


def audit_cboe_vix(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(path)
    frame.columns = [str(c).strip().lower() for c in frame.columns]
    dates = pd.to_datetime(frame["date"], format="%m/%d/%Y", errors="raise")
    numeric = frame[["open", "high", "low", "close"]].apply(pd.to_numeric, errors="coerce")
    valid = numeric[["open", "close"]].min(axis=1).ge(numeric["low"]) & numeric[["open", "close"]].max(axis=1).le(numeric["high"])
    return {
        "rows": len(frame),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "weekend_dates": int((dates.dt.dayofweek >= 5).sum()),
        "missing_ohlc": int(numeric.isna().sum().sum()),
        "invalid_ohlc_rows": int((~valid).sum()),
    }


def period_key_valid(value: Any, frequency: str) -> bool:
    text = str(value).replace(".0", "")
    if not text.isdigit():
        return False
    if frequency == "A":
        return len(text) == 4
    if frequency == "Q":
        return len(text) == 5 and text[-1] in "1234"
    if frequency == "S":
        return len(text) == 5 and text[-1] in "12"
    if frequency == "M":
        return len(text) == 6 and 1 <= int(text[-2:]) <= 12
    return False


def audit_goyal_vintages_zip(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    member_summaries: dict[str, Any] = {}
    totals = Counter()
    with zipfile.ZipFile(path) as archive:
        crc_error = archive.testzip()
        members = sorted(name for name in archive.namelist() if name.lower().endswith(".csv"))
        for name in members:
            frequency = Path(name).stem.rsplit("_", 1)[-1]
            frame = pd.read_csv(archive.open(name), low_memory=False)
            row_labels = frame.iloc[:, 0]
            column_labels = pd.Series(frame.columns[1:])
            numeric = frame.iloc[:, 1:].apply(pd.to_numeric, errors="coerce")
            array = numeric.to_numpy(dtype=float)
            nonfinite = int((~pd.isna(array) & ~pd.DataFrame(array).map(math.isfinite).to_numpy()).sum())
            invalid_rows = sum(not period_key_valid(v, frequency) for v in row_labels)
            invalid_columns = sum(not period_key_valid(v, frequency) for v in column_labels)
            summary = {
                "frequency": frequency,
                "rows": len(frame),
                "vintages": len(frame.columns) - 1,
                "row_start": str(row_labels.iloc[0]).replace(".0", ""),
                "row_end": str(row_labels.iloc[-1]).replace(".0", ""),
                "vintage_start": str(column_labels.iloc[0]),
                "vintage_end": str(column_labels.iloc[-1]),
                "duplicate_rows": int(row_labels.duplicated().sum()),
                "duplicate_vintages": int(column_labels.duplicated().sum()),
                "invalid_period_rows": int(invalid_rows),
                "invalid_period_vintages": int(invalid_columns),
                "missing_values": int(numeric.isna().sum().sum()),
                "infinite_values": nonfinite,
            }
            member_summaries[name] = summary
            for key in ["duplicate_rows", "duplicate_vintages", "invalid_period_rows", "invalid_period_vintages", "infinite_values"]:
                totals[key] += summary[key]
    return {
        "zip_crc_error": crc_error,
        "csv_members": len(member_summaries),
        "total_quality_issues": dict(totals),
        "members": member_summaries,
    }


def audit_goyal_combined(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    workbook = pd.ExcelFile(path)
    summary: dict[str, Any] = {"sheets": workbook.sheet_names, "data_sheets": {}}
    for sheet, key in [("Monthly", "yyyymm"), ("Quarterly", "yyyyq"), ("Annual", "yyyy")]:
        frame = pd.read_excel(path, sheet_name=sheet)
        periods = frame[key].astype(str).str.replace(".0", "", regex=False)
        numeric = frame.drop(columns=key).apply(pd.to_numeric, errors="coerce")
        array = numeric.to_numpy(dtype=float)
        coverage = {}
        for column in numeric.columns:
            mask = numeric[column].notna()
            coverage[column] = {
                "nonmissing": int(mask.sum()),
                "start": periods[mask].iloc[0] if mask.any() else None,
                "end": periods[mask].iloc[-1] if mask.any() else None,
            }
        summary["data_sheets"][sheet] = {
            "rows": len(frame),
            "columns": len(frame.columns),
            "start": periods.iloc[0],
            "end": periods.iloc[-1],
            "duplicate_periods": int(periods.duplicated().sum()),
            "missing_values": int(numeric.isna().sum().sum()),
            "infinite_values": int((~pd.isna(array) & ~pd.DataFrame(array).map(math.isfinite).to_numpy()).sum()),
            "coverage": coverage,
        }
    return summary


def parse_aqr_date(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value), "%m/%d/%Y")
    except ValueError:
        return None


def audit_aqr_daily(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    data_sheets = ["BAB Factors", "MKT", "SMB", "HML FF", "HML Devil", "UMD", "RF"]
    summaries: dict[str, Any] = {}
    for name in data_sheets:
        sheet = workbook[name]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=19, max_row=19))]
        used = max(i for i, value in enumerate(headers) if value is not None) + 1
        headers = headers[:used]
        dates: list[datetime] = []
        duplicate_count = 0
        seen: set[datetime] = set()
        nonmissing = Counter()
        missing = Counter()
        infinite = 0
        for row in sheet.iter_rows(min_row=20, values_only=True):
            date = parse_aqr_date(row[0])
            if date is None:
                continue
            if date in seen:
                duplicate_count += 1
            seen.add(date)
            dates.append(date)
            for column, value in zip(headers[1:], row[1:used]):
                if value is None:
                    missing[str(column)] += 1
                else:
                    nonmissing[str(column)] += 1
                    if isinstance(value, (int, float)) and not math.isfinite(float(value)):
                        infinite += 1
        if not dates:
            raise ValueError(f"aucune donnée datée dans {name}")
        usa_key = "USA" if "USA" in headers else "Risk Free Rate"
        summaries[name] = {
            "rows": len(dates),
            "columns": len(headers),
            "start": min(dates).date().isoformat(),
            "end": max(dates).date().isoformat(),
            "duplicate_dates": duplicate_count,
            "weekend_dates": sum(d.weekday() >= 5 for d in dates),
            "infinite_values": infinite,
            "key_column": usa_key,
            "key_nonmissing": int(nonmissing[usa_key]),
            "key_missing": int(missing[usa_key]),
        }
    return {"sheets": workbook.sheetnames, "data_sheets": summaries}


def read_ssga_sheet(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path, header=3)
    frame = frame.dropna(how="all")
    return frame


def parse_ssga_numeric(values: pd.Series) -> pd.Series:
    """Normalise les séparateurs d'affichage des classeurs officiels State Street."""
    return pd.to_numeric(
        values.astype(str)
        .str.replace("$", "", regex=False)
        .str.replace(",", "", regex=False)
        .str.replace("%", "", regex=False)
        .str.strip()
        .replace({"nan": None, "": None}),
        errors="coerce",
    )


def audit_ssga_nav(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = read_ssga_sheet(path)
    dates = pd.to_datetime(frame["Date"], errors="coerce")
    non_date_rows = int(dates.isna().sum())
    frame = frame.loc[dates.notna()].copy()
    dates = dates.loc[dates.notna()]
    numeric = frame[["NAV", "Shares Outstanding", "Total Net Assets"]].apply(parse_ssga_numeric)
    implied = numeric["NAV"] * numeric["Shares Outstanding"]
    relative_gap = ((implied - numeric["Total Net Assets"]).abs() / numeric["Total Net Assets"]).dropna()
    return {
        "rows": len(frame),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "non_date_footer_rows": non_date_rows,
        "missing_values": int(numeric.isna().sum().sum()),
        "nonpositive_values": int((numeric <= 0).sum().sum()),
        "max_nav_assets_relative_gap": float(relative_gap.max()),
    }


def audit_ssga_premium(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = read_ssga_sheet(path)
    dates = pd.to_datetime(frame["Date"], errors="coerce")
    non_date_rows = int(dates.isna().sum())
    frame = frame.loc[dates.notna()].copy()
    dates = dates.loc[dates.notna()]
    values = parse_ssga_numeric(frame["Premium/Discount"])
    return {
        "rows": len(frame),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "non_date_footer_rows": non_date_rows,
        "missing_values": int(values.isna().sum()),
        "infinite_values": int((values.notna() & ~values.map(math.isfinite)).sum()),
    }


def audit_ssga_distributions(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_excel(path, sheet_name="dividend")
    subset = frame[frame["TICKER"].isin(["SPY", "BIL"])].copy()
    ex_dates = pd.to_datetime(subset["EX-DATE"], errors="raise")
    record_dates = pd.to_datetime(subset["RECORD DATE"], errors="raise")
    payable_dates = pd.to_datetime(subset["PAYABLE DATE"], errors="raise")
    dividends = pd.to_numeric(subset["DIVIDEND ($)"], errors="coerce")
    short_gains = pd.to_numeric(subset["SHORT TERM CAPITAL GAIN ($)"], errors="coerce").fillna(0)
    long_gains = pd.to_numeric(subset["LONG TERM CAPITAL GAIN ($)"], errors="coerce").fillna(0)
    by_ticker = {}
    for ticker in ["SPY", "BIL"]:
        mask = subset["TICKER"].eq(ticker)
        dates = ex_dates[mask]
        by_ticker[ticker] = {
            "rows": int(mask.sum()),
            "start": dates.min().date().isoformat(),
            "end": dates.max().date().isoformat(),
            "duplicate_ex_dates": int(dates.duplicated().sum()),
            "missing_dividends": int(dividends[mask].isna().sum()),
            "negative_dividends": int((dividends[mask] < 0).sum()),
            "nonzero_short_term_capital_gains": int((short_gains[mask] != 0).sum()),
            "nonzero_long_term_capital_gains": int((long_gains[mask] != 0).sum()),
            "record_before_ex_date": int((record_dates[mask] < dates).sum()),
            "payment_before_ex_date": int((payable_dates[mask] < dates).sum()),
        }
    return {"rows_all_funds": len(frame), "rows_spy_bil": len(subset), "by_ticker": by_ticker}


def audit_nyfed_sofrai(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(path)
    dates = pd.to_datetime(frame["Effective Date"], format="%m/%d/%Y", errors="raise")
    fields = ["30-Day Average SOFR", "90-Day Average SOFR", "180-Day Average SOFR", "SOFR Index"]
    numeric = frame[fields].apply(pd.to_numeric, errors="coerce")
    ascending = pd.DataFrame({"date": dates, "index": numeric["SOFR Index"]}).sort_values("date")
    index_nonmissing = ascending.dropna()
    return {
        "rows": len(frame),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "rate_type_values": sorted(frame["Rate Type"].dropna().astype(str).unique().tolist()),
        "nonmissing": {column: int(numeric[column].notna().sum()) for column in fields},
        "infinite_values": int((numeric.notna() & ~numeric.map(math.isfinite)).sum().sum()),
        "nonpositive_index": int((numeric["SOFR Index"].dropna() <= 0).sum()),
        "index_decreases": int((index_nonmissing["index"].diff().dropna() < 0).sum()),
    }


def audit_ofr_fsi(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(path)
    dates = pd.to_datetime(frame["Date"], format="%Y-%m-%d", errors="raise")
    numeric = frame.drop(columns="Date").apply(pd.to_numeric, errors="coerce")
    return {
        "rows": len(frame),
        "columns": list(frame.columns),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "weekend_dates": int((dates.dt.dayofweek >= 5).sum()),
        "missing_values": int(numeric.isna().sum().sum()),
        "infinite_values": int((numeric.notna() & ~numeric.map(math.isfinite)).sum().sum()),
    }


def audit_nfci(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(path, na_values=["."])
    dates = pd.to_datetime(frame["observation_date"], format="%Y-%m-%d", errors="raise")
    numeric = frame[["NFCI", "ANFCI"]].apply(pd.to_numeric, errors="coerce")
    gaps = dates.sort_values().diff().dropna().dt.days
    return {
        "rows": len(frame),
        "start": dates.min().date().isoformat(),
        "end": dates.max().date().isoformat(),
        "duplicate_dates": int(dates.duplicated().sum()),
        "non_friday_dates": int((dates.dt.dayofweek != 4).sum()),
        "gap_days": sorted(int(v) for v in gaps.unique()),
        "missing_values": int(numeric.isna().sum().sum()),
        "infinite_values": int((numeric.notna() & ~numeric.map(math.isfinite)).sum().sum()),
    }


def audit_workbook_structure(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return {
        "sheets": workbook.sheetnames,
        "dimensions": {
            sheet.title: {"rows": sheet.max_row, "columns": sheet.max_column}
            for sheet in workbook.worksheets
        },
    }


def audit_derived_total_returns(path: Path, _: dict[str, Any]) -> dict[str, Any]:
    frame = pd.read_csv(
        path,
        parse_dates=["date", "distribution_record_date", "distribution_payable_date"],
    )
    required = [
        "nav_usd",
        "split_factor_current_per_previous",
        "total_distribution_usd_per_current_share",
        "total_return_nav",
        "growth_factor",
        "total_return_index_base_100",
    ]
    numeric = frame[required].apply(pd.to_numeric, errors="coerce")
    returns = numeric["total_return_nav"]
    growth_error = (numeric["growth_factor"] - (1 + returns)).abs().dropna()
    index_ratio = numeric["total_return_index_base_100"] / numeric["total_return_index_base_100"].shift(1)
    index_error = (index_ratio - numeric["growth_factor"]).abs().dropna()
    split_dates = frame.loc[numeric["split_factor_current_per_previous"].ne(1), "date"]
    event_flag = frame["distribution_event"].astype(str).str.lower().eq("true")
    return {
        "rows": len(frame),
        "ticker_values": sorted(frame["ticker"].dropna().astype(str).unique().tolist()),
        "start": frame["date"].min().date().isoformat(),
        "end": frame["date"].max().date().isoformat(),
        "duplicate_dates": int(frame["date"].duplicated().sum()),
        "missing_nav": int(numeric["nav_usd"].isna().sum()),
        "nonpositive_nav": int((numeric["nav_usd"] <= 0).sum()),
        "missing_returns": int(returns.isna().sum()),
        "infinite_values": int((numeric.notna() & ~numeric.map(math.isfinite)).sum().sum()),
        "distribution_events": int(event_flag.sum()),
        "split_dates": [date.date().isoformat() for date in split_dates],
        "max_growth_identity_error": float(growth_error.max()) if len(growth_error) else 0.0,
        "max_index_identity_error": float(index_error.max()) if len(index_error) else 0.0,
        "initial_index": float(numeric["total_return_index_base_100"].iloc[0]),
        "ending_index": float(numeric["total_return_index_base_100"].iloc[-1]),
    }


def audit_json_document(path: Path, params: dict[str, Any]) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    required = params.get("required_keys", [])
    return {
        "top_level_keys": sorted(payload.keys()),
        "missing_required_keys": sorted(set(required) - set(payload)),
        "status": payload.get("status", payload.get("series_status")),
        "all_required_assertions_passed": payload.get("all_required_assertions_passed"),
    }


HANDLERS = {
    "ff_daily_zip": audit_ff_daily_zip,
    "nasdaq_json": audit_nasdaq_json,
    "fred_rates_csv": audit_fred_rates,
    "cboe_vix_csv": audit_cboe_vix,
    "goyal_vintages_zip": audit_goyal_vintages_zip,
    "goyal_combined_xlsx": audit_goyal_combined,
    "aqr_daily_xlsx": audit_aqr_daily,
    "ssga_nav_xlsx": audit_ssga_nav,
    "ssga_premium_xlsx": audit_ssga_premium,
    "ssga_distributions_xlsx": audit_ssga_distributions,
    "nyfed_sofrai_csv": audit_nyfed_sofrai,
    "ofr_fsi_csv": audit_ofr_fsi,
    "nfci_csv": audit_nfci,
    "workbook_structure": audit_workbook_structure,
    "derived_total_returns_csv": audit_derived_total_returns,
    "json_document": audit_json_document,
}


def compare_expected(actual: Any, expected: Any, prefix: str = "") -> list[str]:
    errors: list[str] = []
    if isinstance(expected, dict):
        if not isinstance(actual, dict):
            return [f"{prefix}: type attendu dict, obtenu {type(actual).__name__}"]
        for key, value in expected.items():
            child = f"{prefix}.{key}" if prefix else key
            if key not in actual:
                errors.append(f"{child}: absent du résultat")
            else:
                errors.extend(compare_expected(actual[key], value, child))
    elif actual != expected:
        errors.append(f"{prefix}: attendu {expected!r}, obtenu {actual!r}")
    return errors


def verify(manifest_path: Path) -> dict[str, Any]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    records = []
    failures = []
    for dataset in manifest["datasets"]:
        record = {"id": dataset["id"], "status": dataset["status"]}
        local_file = dataset.get("local_file")
        if not local_file:
            record["check"] = "NO_LOCAL_FILE_EXPECTED"
            records.append(record)
            continue
        path = ROOT / local_file
        record["path"] = local_file
        if not path.is_file():
            record["check"] = "MISSING"
            failures.append(f"{dataset['id']}: fichier absent: {local_file}")
            records.append(record)
            continue
        actual_sha = sha256_file(path)
        actual_size = path.stat().st_size
        record.update({"bytes": actual_size, "sha256": actual_sha})
        if actual_sha != dataset["sha256"]:
            failures.append(f"{dataset['id']}: SHA-256 modifié")
        if actual_size != dataset["bytes"]:
            failures.append(f"{dataset['id']}: taille modifiée")
        handler_name = dataset.get("audit", {}).get("handler")
        if handler_name:
            try:
                actual_audit = HANDLERS[handler_name](path, dataset.get("audit", {}).get("params", {}))
                record["audit"] = actual_audit
                failures.extend(
                    f"{dataset['id']}: {message}"
                    for message in compare_expected(actual_audit, dataset.get("audit", {}).get("expected", {}))
                )
            except Exception as exc:  # rapporter précisément le fichier fautif
                failures.append(f"{dataset['id']}: audit impossible: {type(exc).__name__}: {exc}")
        record["check"] = "OK" if not any(item.startswith(f"{dataset['id']}:") for item in failures) else "FAILED"
        records.append(record)
    return {
        "manifest": str(manifest_path.relative_to(ROOT)),
        "checked_at": datetime.now().astimezone().isoformat(),
        "records": records,
        "failures": failures,
        "ok": not failures,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--json", action="store_true", help="imprime le rapport JSON complet")
    args = parser.parse_args()
    manifest_path = args.manifest.resolve()
    result = verify(manifest_path)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for record in result["records"]:
            print(f"{record['check']:>22}  {record['status']:<18}  {record['id']}")
        if result["failures"]:
            print("\nÉchecs:", file=sys.stderr)
            for failure in result["failures"]:
                print(f"- {failure}", file=sys.stderr)
        else:
            print(f"\n{len(result['records'])} entrées contrôlées; aucun échec.")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
